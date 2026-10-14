"""Apply print layout to Pandoc's Word output without changing the PDF sources."""
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def flag(parent, tag, value=None):
    node = parent.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        parent.append(node)
    if value is not None:
        node.set(qn('w:val'), str(value))
    return node


def section_properties(base, landscape=False):
    section = deepcopy(base)
    size = section.find(qn('w:pgSz'))
    if size is None:
        size = OxmlElement('w:pgSz')
        section.append(size)
    size.set(qn('w:w'), str(round(Mm(297 if landscape else 210).twips)))
    size.set(qn('w:h'), str(round(Mm(210 if landscape else 297).twips)))
    size.set(qn('w:orient'), 'landscape' if landscape else 'portrait')
    flag(section, 'w:type', 'nextPage')
    return section


def section_break(before, properties):
    paragraph = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr')
    ppr.append(properties)
    paragraph.append(ppr)
    before.addprevious(paragraph)


def format_docx(path):
    doc = Document(path)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Mm(25)
        section.left_margin = section.right_margin = Mm(25)
    base = section_properties(doc.sections[-1]._sectPr)
    final_section = doc.sections[-1]._sectPr
    final_section.getparent().replace(final_section, deepcopy(base))
    wide_blocks = []
    for table in doc.tables:
        ncol = len(table.columns)
        caption = table._tbl.getprevious()
        caption_text = ''.join(caption.itertext()) if caption is not None else ''
        landscape = ncol >= 8
        # Short main-text tables fit on one page; keep their rows and notes
        # together instead of leaving the final election rows on the next page.
        compact_main = path.stem == 'manuscript' and caption_text.startswith('Table ') and len(table.rows) <= 10
        # The full definitions are already in Table 2's Notes. Short Word headers
        # leave room for intervals while keeping the same A–F column references.
        if 'Table 2.' in caption_text:
            headers = ['Election', 'Election day (exact y)', 'A', 'B', 'C', 'D', 'E', 'F']
            if ncol == 9:  # older two-panel tables
                headers.insert(1, 'Conv.')
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
        table.autofit = False
        width = Mm(247 if landscape else 160).twips
        original = [column.width.twips for column in table.columns]
        widths = [round(width * w / sum(original)) for w in original]
        for column, w in zip(table.columns, widths):
            column.width = Mm(w * 25.4 / 1440)
        for row_index, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            flag(trpr, 'w:cantSplit')
            if row_index == 0:
                flag(trpr, 'w:tblHeader')
            # Panel headings belong to the whole table, rather than a narrow
            # first column. Merge only rows whose other cells are all empty.
            cells = row.cells
            panel = row_index > 0 and cells[0].text.strip() and all(not c.text.strip() for c in cells[1:])
            if panel:
                cells[0].merge(cells[-1])
                cells = [row.cells[0]]
                cells[0].width = Mm(width * 25.4 / 1440)
            else:
                for cell, w in zip(cells, widths):
                    cell.width = Mm(w * 25.4 / 1440)
            for i, cell in enumerate(cells):
                for paragraph in cell.paragraphs:
                    fmt = paragraph.paragraph_format
                    fmt.space_before = Pt(0)
                    fmt.space_after = Pt(2)
                    fmt.line_spacing = 1
                    fmt.keep_with_next = row_index == 0 or bool(panel) or (compact_main and row_index < len(table.rows) - 1)
                    fmt.keep_together = True
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        # Caption stays with the first table row. Notes use table-sized text.
        if caption is not None and caption.tag == qn('w:p'):
            flag(caption.get_or_add_pPr(), 'w:keepNext')
        end = table._tbl
        following = end.getnext()
        if following is not None and following.tag == qn('w:p'):
            from docx.text.paragraph import Paragraph
            notes = Paragraph(following, table._parent)
            if notes.text.startswith('Notes.'):
                if compact_main:
                    for cell in table.rows[-1].cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
                    notes.paragraph_format.keep_together = True
                notes.paragraph_format.line_spacing = 1
                notes.paragraph_format.space_before = Pt(4)
                notes.paragraph_format.space_after = Pt(6)
                for run in notes.runs:
                    run.font.size = Pt(9)
                end = following
        if landscape:
            if wide_blocks and wide_blocks[-1][1].getnext() is caption:
                wide_blocks[-1] = (wide_blocks[-1][0], end)
            else:
                wide_blocks.append((caption, end))
    for caption, end in wide_blocks:
        section_break(caption, section_properties(base))
        following = end.getnext()
        if following is not None:
            section_break(following, section_properties(base, landscape=True))
    # Keep each image on the same page as its numbered caption.
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_with_next = True
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.save(path)
